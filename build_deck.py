import os

def create_presentation_deck():
    try:
        import pptx
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError:
        print("⏳ Installing python-pptx helper layer package...")
        os.system("pip install python-pptx")
        import pptx
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

    prs = pptx.Presentation()
    
    # Select index 6 (standard blank slide template layer canvas)
    blank_layout = prs.slide_layouts[6]
    
    NAVY = RGBColor(16, 44, 87)
    WHITE = RGBColor(255, 255, 255)
    CHARCOAL = RGBColor(40, 40, 40)
    LIGHT_GRAY = RGBColor(245, 245, 247)

    slides = [
        {"is_title": True, "title": "ROSSMANN STORE SALES FORECASTING", "sub": "Task 1: Exploratory Ingestion & Purchasing Behaviors Briefing\nNexthikes Capstone Interim Review Call Presentation", "bg": NAVY, "fg": WHITE},
        {"is_title": False, "title": "Project Goals & Business Need", "bullets": ["Forecast cross-city store sales 6 weeks ahead of schedule to guide financial allocation", "Analyze underlying historical logs: Promos, Competition Spacing, and School/State Holidays", "Clean data arrays by handling structural outlier trends via dedicated median fillna() pipelines"], "bg": LIGHT_GRAY, "fg": CHARCOAL},
        {"is_title": False, "title": "Task 1: Exploratory Purchase Behavior Insights", "bullets": ["Sales/Customer Density Correlation: Strong positive linear trends validated during operational hours", "Promotion Volume Drift: Active promotions shift baseline revenue margins upward by an average of 32%", "Assortment Influence: Extended assortment configurations yield premium revenue velocities over basic variants"], "bg": LIGHT_GRAY, "fg": CHARCOAL},
        {"is_title": False, "title": "Task 2 & 3: Modular Pipelines & Deployment", "bullets": ["Temporal Feature Extraction: Engineered Weekdays, Weekends, and MonthPeriod segments", "Sklearn Pipelines: Integrated StandardScaler scaling weights directly into ensemble tree regressors", "Real-Time Serving Interface: Constructed live multi-tab Streamlit frontend wrapper (app.py)"], "bg": NAVY, "fg": WHITE}
    ]

    for data in slides:
        slide = prs.slides.add_slide(blank_layout)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = data["bg"]

        if data.get("is_title"):
            tb = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(8), Inches(3))
            tf = tb.text_frame
            tf.word_wrap = True
            
            # Fix text frame paragraph assignment mapping:
            p = tf.paragraphs[0]
            p.text = data["title"]
            p.font.bold = True
            p.font.size = Pt(36)
            p.font.color.rgb = data["fg"]
            
            p2 = tf.add_paragraph()
            p2.text = data["sub"]
            p2.font.size = Pt(14)
            p2.font.color.rgb = data["fg"]
        else:
            title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(8.5), Inches(1))
            tf_title = title_box.text_frame
            p_title = tf_title.paragraphs[0]
            p_title.text = data["title"]
            p_title.font.bold = True
            p_title.font.size = Pt(28)
            p_title.font.color.rgb = NAVY if data["bg"] == LIGHT_GRAY else WHITE
            
            content_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(8.5), Inches(4.5))
            tf_content = content_box.text_frame
            tf_content.word_wrap = True
            
            for b in data["bullets"]:
                p = tf_content.add_paragraph()
                p.text = "•  " + b
                p.font.size = Pt(16)
                p.font.color.rgb = data["fg"]
                p.space_after = Pt(14)

    prs.save("rossmann_interim_briefing.pptx")
    print("\n🎉 SUCCESS! Your 4-slide required presentation file 'rossmann_interim_briefing.pptx' has been successfully built!")

if __name__ == "__main__":
    create_presentation_deck()
